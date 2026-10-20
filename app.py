import streamlit as st
from google import genai
import json, io
import pdfplumber
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. 페이지 기본 설정
st.set_page_config(page_title="AI 시험 문제 출제기", page_icon="📝", layout="wide")
st.title("📝 AI 맞춤형 시험 문제 출제기")
st.caption("PDF/한글 파일 또는 텍스트를 입력받아 AI가 난이도별 시험지를 생성합니다.")

# 2. API 키 불러오기 (비밀 키 금고에서 안전하게 호출)
API_KEY = st.secrets.get("GEMINI_API_KEY", "")
if not API_KEY:
    st.error("설정에서 GEMINI_API_KEY를 등록해주세요!")
    st.stop()

# 3. 파일에서 텍스트 추출하는 함수
def extract_text(file):
    name = file.name.lower()
    text = ""
    try:
        # PDF 파일 읽기
        if name.endswith(".pdf"):
            with pdfplumber.open(file) as pdf:
                for page in pdf.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
        # 텍스트 파일 읽기
        elif name.endswith(".txt"):
            text = file.read().decode("utf-8", errors="ignore")
        # 한글(HWPX) 파일 읽기
        elif name.endswith(".hwpx"):
            import zipfile
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(file) as z:
                xml_files = [f for f in z.namelist() if f.startswith("Contents/section") and f.endswith(".xml")]
                for xf in xml_files:
                    tree = ET.fromstring(z.read(xf))
                    for node in tree.iter():
                        if node.text:
                            text += node.text + " "
        # 구버전 한글(HWP) 파일 읽기
        elif name.endswith(".hwp"):
            try:
                import olefile
                f_bytes = file.read()
                ole = olefile.OleFileIO(f_bytes)
                if ole.exists('PrvText'):
                    stream = ole.openstream('PrvText').read()
                    text = stream.decode('utf-16le', errors='ignore')
                else:
                    st.warning("이 HWP 파일은 텍스트 미리보기가 포함되어 있지 않습니다. HWPX나 PDF로 저장 후 올려주세요.")
            except Exception as e:
                st.warning("HWP 텍스트 추출 중 문제가 발생했습니다. PDF나 HWPX로 변환하여 올려주세요.")
    except Exception as e:
        st.error(f"파일을 읽는 도중 오류가 발생했습니다: {e}")
    return text.strip()

# 4. Word 다운로드 문서 생성 함수
def create_word(questions, passage):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.8)

    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run("【 맞 춤 형 평 가 문 항 】")
    r.font.name = "맑은 고딕"; r.font.size = Pt(18); r.font.bold = True
    doc.add_paragraph()

    bp = doc.add_paragraph()
    br = bp.add_run("※ [다음 글을 읽고 물음에 답하시오.]")
    br.font.bold = True; br.font.size = Pt(11)

    pp = doc.add_paragraph()
    pp.paragraph_format.left_indent = Inches(0.2)
    pp.paragraph_format.right_indent = Inches(0.2)
    pr = pp.add_run(passage[:2000] + ("..." if len(passage) > 2000 else ""))
    pr.font.size = Pt(9.5)
    pr.font.color.rgb = RGBColor(70, 70, 70)
    doc.add_paragraph("―" * 45)

    for q in questions:
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(8)
        qr = qp.add_run(f"[{q.get('number', '')}] {q.get('question', '')} ")
        qr.font.bold = True; qr.font.size = Pt(11)
        dr = qp.add_run(f"(난이도: {q.get('difficulty', '보통')})")
        dr.font.size = Pt(9); dr.font.color.rgb = RGBColor(120, 120, 120)

        for opt in q.get("options", []):
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Inches(0.25)
            op.paragraph_format.space_after = Pt(2)
            opr = op.add_run(opt)
            opr.font.size = Pt(10)

    doc.add_page_break()
    ap = doc.add_paragraph()
    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ar = ap.add_run("【 정 답  및  해 설 】")
    ar.font.bold = True; ar.font.size = Pt(16)
    doc.add_paragraph()

    for q in questions:
        ip = doc.add_paragraph()
        ip.paragraph_format.space_before = Pt(6)
        nr = ip.add_run(f"{q.get('number', '')}번 정답: {q.get('answer', '')}\n")
        nr.font.bold = True; nr.font.size = Pt(11)
        nr.font.color.rgb = RGBColor(30, 80, 180)
        er = ip.add_run(f"• 해설: {q.get('explanation', '')}")
        er.font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

# 5. UI 및 입력 영역
with st.sidebar:
    st.header("⚙️ 출제 설정")
    cnt = st.number_input("전체 문항 수", min_value=1, max_value=20, value=3, step=1)
    diff = st.selectbox("난이도 구성", ["난이도 골고루 (상/중/하 혼합)", "상 (심화/추론 중심)", "중 (이해/적용 중심)", "하 (기본/사실 확인)"])

# 파일 업로드 or 직접 입력
uploaded_file = st.file_uploader("📂 지문/학습 자료 파일 업로드 (PDF, HWP, HWPX, TXT)", type=["pdf", "hwp", "hwpx", "txt"])

extracted_content = ""
if uploaded_file:
    with st.spinner("파일에서 텍스트를 추출하고 있습니다..."):
        extracted_content = extract_text(uploaded_file)
        if extracted_content:
            st.success(f"✅ '{uploaded_file.name}'에서 텍스트를 성공적으로 추출했습니다! (글자 수: {len(extracted_content):,}자)")

passage = st.text_area(
    "📌 분석할 지문 내용 (직접 수정하거나 입력 가능)",
    value=extracted_content,
    height=260,
    placeholder="파일을 올리면 여기에 내용이 자동으로 채워지며, 직접 텍스트를 붙여넣으셔도 됩니다."
)

if "questions" not in st.session_state:
    st.session_state.questions = None
if "passage" not in st.session_state:
    st.session_state.passage = ""

# 6. 문제 출제 버튼
if st.button("🚀 시험 문제 생성하기", type="primary"):
    if not passage.strip():
        st.error("지문 내용이 비어있습니다. 파일을 올리거나 글을 입력해주세요!")
    else:
        with st.spinner("AI가 지문을 정밀 분석하여 문제를 출제하고 있습니다 (약 15~20초)..."):
            try:
                client = genai.Client(api_key=API_KEY)
                prompt = f"""
당신은 시험 출제 전문가입니다. 아래 [지문]을 정밀 분석하여 총 {cnt}개의 객관식(5지선다) 문제를 만들어주세요.

[출제 요구사항]
- 난이도 기준: {diff}
- 각 문제마다 난이도(상/중/하), 문제 번호, 정답, 상세한 정답 및 오답 해설을 작성하세요.
- 반드시 아래 JSON 스키마 규격으로만 응답하세요.

{{
  "questions": [
    {{
      "number": 1,
      "difficulty": "중",
      "question": "문제 내용",
      "options": ["① 보기 1", "② 보기 2", "③ 보기 3", "④ 보기 4", "⑤ 보기 5"],
      "answer": "① 보기 1",
      "explanation": "해설 내용"
    }}
  ]
}}

[지문]
{passage}
"""
                res = client.models.generate_content(
                    model='gemini-3.8-flash',
                    contents=prompt,
                    config={'response_mime_type': 'application/json'}
                )
                data = json.loads(res.text)
                st.session_state.questions = data.get("questions", [])
                st.session_state.passage = passage
                st.success("🎉 문제가 성공적으로 생성되었습니다!")
            except Exception as e:
                st.error(f"오류가 발생했습니다: {e}")

# 7. 문제 출력 및 다운로드
if st.session_state.questions:
    w_data = create_word(st.session_state.questions, st.session_state.passage)
    st.download_button(
        label="📥 Word 시험지 & 정답지 다운로드 (.docx)",
        data=w_data,
        file_name="AI_시험지_및_정답해설.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    st.markdown("---")
    for q in st.session_state.questions:
        col1, col2 = st.columns([1, 8])
        with col1:
            st.info(f"**난이도: {q.get('difficulty', '보통')}**")
        with col2:
            st.subheader(f"Q{q.get('number', '')}. {q.get('question', '')}")
        for opt in q.get("options", []):
            st.write(opt)
        with st.expander("🔍 정답 및 해설 보기"):
            st.markdown(f"**정답:** {q.get('answer', '')}")
            st.markdown(f"**해설:** {q.get('explanation', '')}")
        st.markdown("---")
